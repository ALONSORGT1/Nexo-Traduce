import base64
import io
import os
import unittest
import zipfile
from types import SimpleNamespace
from unittest.mock import Mock, patch
from PIL import Image
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject
from backend.app import create_app
from backend.files import DocumentExtractor
from backend.services import TranslationService
from backend.errors import AppError


def png():
    output = io.BytesIO(); Image.new('RGB', (100, 100), 'white').save(output, 'PNG'); return output.getvalue()


def pdf():
    writer = PdfWriter(); page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 50 700 Td (First section. Meeting tomorrow.) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO(); writer.write(output); return output.getvalue()


class HttpTests(unittest.TestCase):
    def setUp(self):
        self.env = patch.dict(os.environ, {'APP_ACCESS_CODE': '', 'ALLOWED_ORIGINS': 'http://localhost:5500'})
        self.env.start(); self.addCleanup(self.env.stop)
        self.service = Mock()
        self.service.translate.side_effect = lambda text, source, target, *args: {'original': text, 'translation': 'Translated content', 'source': source, 'target': target}
        self.client = create_app(self.service).test_client()

    def test_chat_both_directions(self):
        for source, target in [('es', 'en'), ('en', 'es')]:
            response = self.client.post('/api/chat', json={'source': source, 'target': target, 'text': 'Hola'} )
            self.assertEqual(response.status_code, 200); self.assertEqual(response.json['source'], source)

    def test_invalid_chat(self):
        for data in [[], None, {}, {'text': ' '}, {'source': 'es', 'target': 'es', 'text': 'Hi'}, {'source': 'es', 'target': 'en', 'text': 'Hi', 'history': [None]}]:
            with self.subTest(data=data):
                self.assertIn(self.client.post('/api/chat', json=data).status_code, [400, 415])
        self.service.translate.assert_not_called()

    def test_cors_and_options(self):
        self.assertEqual(self.client.options('/api/chat', headers={'Origin': 'https://evil.test'}).status_code, 403)
        response = self.client.options('/api/chat', headers={'Origin': 'http://localhost:5500'})
        self.assertEqual(response.status_code, 204)
        self.assertEqual(response.headers['Access-Control-Allow-Origin'], 'http://localhost:5500')

    def test_access_code(self):
        with patch.dict(os.environ, {'APP_ACCESS_CODE': 'example-code'}):
            payload = {'text': 'Hi', 'source': 'en', 'target': 'es'}
            self.assertEqual(self.client.post('/api/chat', json=payload).status_code, 401)
            self.assertEqual(self.client.post('/api/chat', json=payload, headers={'X-Access-Code': 'example-code'}).status_code, 200)

    def test_private_paths_are_not_served(self):
        for path in ['/.env.local', '/backend/app.py', '/.git/config', '/assets/../../.env.local']:
            self.assertEqual(self.client.get(path).status_code, 404)

    def upload(self, path, data, filename):
        return self.client.post('/api/' + path, data={'file': (io.BytesIO(data), filename), 'source': 'es', 'target': 'en'})

    def test_text_document(self):
        response = self.upload('documents', 'Título\n\nPrimer párrafo.\n\nSegunda sección.'.encode(), 'sample.txt')
        self.assertEqual(response.status_code, 200); self.assertIn('Segunda', response.json['sections'][0]['original'])

    def test_docx_and_pdf(self):
        doc = Document(); doc.add_heading('Heading', 1); doc.add_paragraph('First paragraph.')
        table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = 'Name'; table.cell(0, 1).text = 'Value'
        doc.add_paragraph('Last paragraph.'); output = io.BytesIO(); doc.save(output)
        response = self.upload('documents', output.getvalue(), 'sample.docx')
        self.assertEqual(response.status_code, 200)
        text = response.json['sections'][0]['original']; self.assertLess(text.index('Name'), text.index('Last'))
        self.assertEqual(self.upload('documents', pdf(), 'sample.pdf').status_code, 200)

    def test_invalid_documents(self):
        for raw, filename in [(b'', 'empty.txt'), (b'fake', 'fake.pdf'), (b'fake', 'fake.docx'), (b'\xff', 'bad.txt'), (b'\0', 'bad.txt'), (b'hello', 'old.doc')]:
            with self.subTest(filename=filename): self.assertEqual(self.upload('documents', raw, filename).status_code, 400)
        self.assertEqual(self.upload('documents', b'a' * (3 * 1024 * 1024 + 1), 'huge.txt').status_code, 413)
        self.assertEqual(self.upload('documents', b'a' * 24001, 'long.txt').status_code, 413)

    def test_blank_pdf(self):
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100); output = io.BytesIO(); writer.write(output)
        self.assertEqual(self.upload('documents', output.getvalue(), 'scan.pdf').status_code, 400)

    def test_forged_audio_and_image(self):
        self.assertEqual(self.upload('transcribe', b'not audio', 'fake.mp3').status_code, 400)
        self.assertEqual(self.upload('images', b'not image', 'fake.png').status_code, 400)
        self.assertEqual(self.upload('images', png(), 'fake.jpg').status_code, 400)

    def test_image_route(self):
        self.service.image_translation.return_value = {'original': 'Hola', 'translation': 'Hello', 'source': 'es', 'target': 'en'}
        self.assertEqual(self.upload('images', png(), 'test.png').status_code, 200)
        self.service.image_translation.assert_called_once()

    def test_errors_hide_internal_details(self):
        self.service.translate.side_effect = RuntimeError('secret internal detail')
        response = self.client.post('/api/chat', json={'text': 'Hi', 'source': 'en', 'target': 'es'})
        self.assertEqual(response.status_code, 500); self.assertNotIn('secret', response.get_data(as_text=True))

    def test_missing_file_and_oversized_body(self):
        self.assertEqual(self.client.post('/api/documents', data={'source': 'es', 'target': 'en'}).status_code, 400)
        self.assertEqual(self.client.post('/api/chat', data=b'a' * (4 * 1024 * 1024 + 1), content_type='application/json').status_code, 413)


class ServiceTests(unittest.TestCase):
    def test_translation_does_not_store_and_context_is_data(self):
        client = Mock(); client.responses.create.return_value = SimpleNamespace(output_text='Hello', status='completed')
        service = TranslationService(client); self.assertEqual(service.translate('Hola', 'es', 'en')['translation'], 'Hello')
        args = client.responses.create.call_args.kwargs; self.assertFalse(args['store']); self.assertIn('never', args['instructions'].lower())

    def test_incomplete_response(self):
        with self.assertRaises(AppError): TranslationService._output(SimpleNamespace(output_text='Partial', status='incomplete'))

    def test_unreadable_and_invalid_image_output(self):
        for output in ['{}', 'not json', '{"readable":false,"original":"","translation":"","warning":""}']:
            client = Mock(); client.responses.create.return_value = SimpleNamespace(output_text=output, status='completed')
            with self.assertRaises(AppError): TranslationService(client).image_translation(png(), 'es', 'en')

    def test_silent_audio(self):
        client = Mock(); client.audio.transcriptions.create.return_value = SimpleNamespace(text=' ')
        with self.assertRaises(AppError): TranslationService(client).transcribe(b'audio', '.wav', 'es')

    def test_speech_and_image_contract(self):
        client = Mock(); client.audio.speech.create.return_value = SimpleNamespace(content=b'mp3')
        service = TranslationService(client); self.assertTrue(service.speech('Hola')['audio'].startswith('data:audio/mpeg;'))
        client.images.edit.return_value = SimpleNamespace(data=[SimpleNamespace(b64_json=base64.b64encode(png()).decode())])
        self.assertTrue(service.edit_image(png(), 'en', 'Hola', 'Hello')['image'].startswith('data:image/jpeg;'))
        self.assertEqual(client.images.edit.call_args.kwargs['image'][0], 'original.jpg')

    def test_sections_preserve_content(self):
        text = '\n\n'.join(['Paragraph ' + str(i) + ' ' + 'x' * 1300 for i in range(12)])
        parts = DocumentExtractor().extract(text.encode(), '.txt')
        self.assertGreater(len(parts), 1); self.assertEqual('\n\n'.join(parts), text)


if __name__ == '__main__': unittest.main()
